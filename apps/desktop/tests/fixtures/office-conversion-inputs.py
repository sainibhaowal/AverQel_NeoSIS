"""Create one-page Office inputs with the packaged authoring runtime."""
from pathlib import Path
import sys
from zipfile import ZIP_DEFLATED, ZIP_STORED, ZipFile
from docx import Document
from openpyxl import Workbook
from pptx import Presentation
from pptx.util import Inches

root = Path(sys.argv[1])
document = Document()
document.add_paragraph("Desktop Office conversion 42")
document.save(root / "input.docx")
workbook = Workbook()
workbook.active.append(["Desktop Office conversion", 42])
workbook.save(root / "input.xlsx")
workbook.close()
presentation = Presentation()
slide = presentation.slides.add_slide(presentation.slide_layouts[6])
slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1)).text = "Desktop Office conversion 42"
presentation.save(root / "input.pptx")


def write_odf(path, mimetype, body):
    with ZipFile(path, "w") as archive:
        archive.writestr("mimetype", mimetype, compress_type=ZIP_STORED)
        archive.writestr("content.xml", f'''<?xml version="1.0" encoding="UTF-8"?>
<office:document-content xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0"
 xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0"
 xmlns:table="urn:oasis:names:tc:opendocument:xmlns:table:1.0"
 xmlns:draw="urn:oasis:names:tc:opendocument:xmlns:drawing:1.0" office:version="1.3">
 <office:body>{body}</office:body>
</office:document-content>''', compress_type=ZIP_DEFLATED)
        archive.writestr("META-INF/manifest.xml", f'''<?xml version="1.0" encoding="UTF-8"?>
<manifest:manifest xmlns:manifest="urn:oasis:names:tc:opendocument:xmlns:manifest:1.0" manifest:version="1.3">
 <manifest:file-entry manifest:full-path="/" manifest:media-type="{mimetype}"/>
 <manifest:file-entry manifest:full-path="content.xml" manifest:media-type="text/xml"/>
</manifest:manifest>''', compress_type=ZIP_DEFLATED)


write_odf(root / "input.odt", "application/vnd.oasis.opendocument.text",
          "<office:text><text:p>Desktop Office conversion ODT</text:p></office:text>")
write_odf(root / "input.ods", "application/vnd.oasis.opendocument.spreadsheet",
          "<office:spreadsheet><table:table table:name=\"Sheet1\"><table:table-row>"
          "<table:table-cell office:value-type=\"string\"><text:p>Desktop Office conversion ODS</text:p>"
          "</table:table-cell></table:table-row></table:table></office:spreadsheet>")
write_odf(root / "input.odp", "application/vnd.oasis.opendocument.presentation",
          "<office:presentation><draw:page draw:name=\"Slide 1\"/></office:presentation>")
(root / "input.rtf").write_text(r"{\rtf1\ansi Desktop Office conversion RTF\par}", encoding="ascii")
